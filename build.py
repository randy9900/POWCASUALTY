import datetime
import json
import pathlib
import shutil
import sys
import openpyxl

root = pathlib.Path(__file__).resolve().parent
source = pathlib.Path(sys.argv[1]).resolve()
book = openpyxl.load_workbook(source, read_only=True, data_only=True)
def convert(v):
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v.isoformat()
    return v
sheets = []
for sheet in book:
    values = list(sheet.values)
    grid = sheet.title == 'Junyo Merge Log'
    headers = [openpyxl.utils.get_column_letter(i+1) for i in range(sheet.max_column)] if grid else [str(v) if v is not None else f'Unnamed column {openpyxl.utils.get_column_letter(i+1)}' for i,v in enumerate(values[0])]
    rows = [[convert(v) for v in r] for r in (values if grid else values[1:]) if grid or any(v is not None for v in r)]
    sheets.append({'name': sheet.title, 'headers': headers, 'rows': rows})
data = json.dumps(sheets, ensure_ascii=False, allow_nan=False).replace('<','\\u003c').replace('>','\\u003e').replace('&','\\u0026')
(root / 'index.html').write_text((root/'template.html').read_text().replace('__DATA__',data), encoding='utf-8')
target = root / 'Master-Casualty-Roster.xlsx'
if source != target:
    shutil.copyfile(source, target)
(root / '.nojekyll').touch()
print([(s['name'],len(s['rows'])) for s in sheets])
